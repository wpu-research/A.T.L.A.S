"""Build the IEEE Access version of the manuscript.

Takes the official IEEE Access Word template (Access-Template-2024.docx), keeps its
first-page block and section/column layout, and fills it with the content of the
already-built manuscript (../TranscriptionSync_revised.docx), mapping each block to
a template style. Manual numbering from the manuscript is kept, so template
auto-numbering is switched off on every inserted paragraph.
"""
import copy, re, sys
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

HERE = Path(__file__).resolve().parent
SRC = HERE.parent / 'TranscriptionSync_revised.docx'
TPL = HERE / 'Access-Template-2024.docx'
OUT = HERE / 'TranscriptionSync_IEEE_Access.docx'
ABSTRACT = (HERE / 'abstract_ieee.txt').read_text().strip()
INDEX = ('ARKit blendshapes, embodied conversational agents, forced alignment, grapheme-to-phoneme conversion, '
         'large language models, lip synchronization, real-time facial animation, speech-to-speech models, viseme synthesis')

src = Document(SRC)
doc = Document(TPL)
body = doc.element.body


def no_num(p):
    """Disable list numbering inherited from a template style."""
    pPr = p._p.get_or_add_pPr()
    numPr = pPr.find(qn('w:numPr'))
    if numPr is None:
        numPr = OxmlElement('w:numPr'); pPr.insert(0, numPr)
    for c in list(numPr):
        numPr.remove(c)
    il = OxmlElement('w:ilvl'); il.set(qn('w:val'), '0'); numPr.append(il)
    ni = OxmlElement('w:numId'); ni.set(qn('w:val'), '0'); numPr.append(ni)


def set_text(p, text):
    for r in list(p._p.iter(qn('w:r'))):
        r.getparent().remove(r)
    for h in list(p._p.iter(qn('w:hyperlink'))):
        h.getparent().remove(h)
    p.add_run(text)


# ---- 1. first-page block -------------------------------------------------------
paras = doc.paragraphs
by_style = {}
for p in paras:
    by_style.setdefault(p.style.name, []).append(p)
set_text(by_style['Paper Title'][0], 'TranscriptionSync: Training-Free Lip Synchronization for Native-Audio Large Language Models and a Layer-Wise Evaluation Against Video-Measured Articulation')
au = by_style['AU'][0]; set_text(au, 'MURAT ARSLAN')
r = au.add_run('1'); r.font.superscript = True
pis = by_style['PI_No Space'] + by_style['PI']
set_text(pis[0], ''); r = pis[0].add_run('1'); r.font.superscript = True
pis[0].add_run('Department of Software Engineering, Altınbaş Cyprus University, Cyprus')
for p in pis[1:-1]:
    p._p.getparent().remove(p._p)
set_text(pis[-1], 'Corresponding author: Murat Arslan (e-mail: research@wpu.edu.tr).')
fn = by_style['footnote text'][0]
set_text(fn, '')
r = fn.add_run('[Funding statement to be completed by the author, e.g., "This work received no external funding."]')
r.font.highlight_color = 7
ab = by_style['Abstract'][0]; set_text(ab, '')
r = ab.add_run('ABSTRACT '); r.bold = True
ab.add_run(ABSTRACT)
it = by_style['IT'][0]; set_text(it, '')
r = it.add_run('INDEX TERMS '); r.bold = True
it.add_run(INDEX + '.')

# ---- 2. drop the template body after the first section break ---------------------
children = list(body)
first_sect = None
for i, el in enumerate(children):
    if el.tag == qn('w:p') and el.find('.//' + qn('w:sectPr')) is not None:
        first_sect = i; break
final_sect = body.find(qn('w:sectPr'))
for el in children[first_sect + 1:]:
    if el is not final_sect:
        body.remove(el)


def add_par(style, runs_from=None, text=None, keep_num=False):
    p = doc.add_paragraph(style=style)
    body.remove(p._p); body.insert(list(body).index(final_sect), p._p)
    if not keep_num:
        no_num(p)
    if text is not None:
        p.add_run(text)
    if runs_from is not None:
        for r in runs_from.iter(qn('w:r')):
            t = ''.join(x.text or '' for x in r.iter(qn('w:t')))
            if not t:
                continue
            nr = p.add_run(t)
            rPr = r.find(qn('w:rPr'))
            if rPr is not None:
                nr.bold = rPr.find(qn('w:b')) is not None or None
                nr.italic = rPr.find(qn('w:i')) is not None or None
                va = rPr.find(qn('w:vertAlign'))
                if va is not None:
                    if va.get(qn('w:val')) == 'superscript': nr.font.superscript = True
                    if va.get(qn('w:val')) == 'subscript': nr.font.subscript = True
    return p


def image_bytes(p_el, part):
    blip = p_el.find('.//' + qn('a:blip'))
    rid = blip.get(qn('r:embed'))
    return part.related_parts[rid].blob


import io
WIDE = {'fig3.png', 'fig_b2.png', 'fig_e4_summary.png', 'fig_e4_lag.png', 'image2.png', 'fig5.png'}

# ---- 3. map the manuscript body ----------------------------------------------------
started = False
section_title = ''
refs_started = False
src_part = src.part
ref_n = 0
atlas_ref = None
for el in src.element.body:
    if el.tag == qn('w:tbl'):
        if started:
            rows = []
            for tr in el.iter(qn('w:tr')):
                row = []
                for tc in tr.findall(qn('w:tc')):
                    cell_runs = []
                    for r in tc.iter(qn('w:r')):
                        t_ = ''.join(x.text or '' for x in r.iter(qn('w:t')))
                        if t_:
                            rPr = r.find(qn('w:rPr'))
                            cell_runs.append((t_, rPr is not None and rPr.find(qn('w:b')) is not None))
                    row.append(cell_runs)
                rows.append(row)
            ncol = max(len(r) for r in rows)
            nt = doc.add_table(rows=len(rows), cols=ncol)
            body.remove(nt._tbl); body.insert(list(body).index(final_sect), nt._tbl)
            lens = [max(4, max(min(40, sum(len(t_) for t_, _ in r[i])) if i < len(r) else 0 for r in rows)) for i in range(ncol)]
            tot = sum(lens); ws = [int(5000 * l / tot) for l in lens]; ws[-1] = 5000 - sum(ws[:-1])
            tblPr = nt._tbl.tblPr
            b = OxmlElement('w:tblBorders')
            for e in ('top', 'bottom', 'insideH'):
                x = OxmlElement(f'w:{e}'); x.set(qn('w:val'), 'single'); x.set(qn('w:sz'), '4'); x.set(qn('w:color'), '000000'); b.append(x)
            tblPr.append(b)
            tw = OxmlElement('w:tblW'); tw.set(qn('w:w'), '5000'); tw.set(qn('w:type'), 'dxa'); tblPr.append(tw)
            lay = OxmlElement('w:tblLayout'); lay.set(qn('w:type'), 'fixed'); tblPr.append(lay)
            for gc, w in zip(nt._tbl.tblGrid.findall(qn('w:gridCol')), ws):
                gc.set(qn('w:w'), str(w))
            for ri, r in enumerate(rows):
                for ci in range(ncol):
                    cell = nt.cell(ri, ci)
                    tcPr = cell._tc.get_or_add_tcPr()
                    for old in tcPr.findall(qn('w:tcW')): tcPr.remove(old)
                    w = OxmlElement('w:tcW'); w.set(qn('w:w'), str(ws[ci])); w.set(qn('w:type'), 'dxa'); tcPr.append(w)
                    cp = cell.paragraphs[0]
                    cp.paragraph_format.space_before = Pt(0); cp.paragraph_format.space_after = Pt(0)
                    for t_, bold in (r[ci] if ci < len(r) else []):
                        rr = cp.add_run(t_); rr.font.size = Pt(7); rr.bold = bold or (ri == 0)
            spacer = add_par('PARA', text='')
        continue
    if el.tag != qn('w:p'):
        continue
    from docx.text.paragraph import Paragraph
    sp = Paragraph(el, src)
    st = sp.style.name
    txt = sp.text.strip()
    if st == 'SecHead':
        section_title = txt
        if txt.startswith('I. INTRODUCTION') or txt.startswith('I\\. INTRODUCTION'):
            started = True
        if not started:
            continue
        if txt.startswith('REFERENCES'):
            # acknowledgment before references (IEEE order)
            add_par('H1', text='ACKNOWLEDGMENT')
            add_par('PARA', text=('The A.T.L.A.S base code derives from the open-source Atlas-MK37 assistant by FatihMakes '
                                  '{ATLAS}. During the revision, the author used Claude (Anthropic) {AI}, an artificial-intelligence '
                                  'system, to assist in drafting and editing parts of the text, in writing and running evaluation '
                                  'scripts, and in analyzing results. All text, code and results were reviewed by the author, who '
                                  'takes full responsibility for the content of this article.'))
            refs_started = True
        add_par('H1', text=txt.replace('\\.', '.'))
        continue
    if not started:
        continue
    if el.find('.//' + qn('w:drawing')) is not None:
        name = ''
        blip = el.find('.//' + qn('a:blip'))
        rel = src_part.rels[blip.get(qn('r:embed'))]
        name = Path(rel.target_ref).name
        p = add_par('PARA')
        p.alignment = 1
        p.add_run().add_picture(io.BytesIO(image_bytes(el, src_part)), width=Inches(3.4))
        continue
    if not txt:
        continue
    if st == 'CaptionText':
        pc = add_par('Figure Caption', runs_from=el)
        if txt.startswith('TABLE'):
            pc.alignment = 1; pc.paragraph_format.keep_with_next = True
        continue
    if st == 'Equation':
        add_par('Equation', runs_from=el); continue
    if st == 'RefText':
        ref_n += 1
        if 'Atlas-MK37' in txt:
            atlas_ref = ref_n
        add_par('References', runs_from=el)
        continue
    # bold subsection headings "A. Overview"
    runs = [r for r in el.iter(qn('w:r')) if ''.join(x.text or '' for x in r.iter(qn('w:t'))).strip()]
    if len(runs) == 1 and runs[0].find(qn('w:rPr')) is not None and runs[0].find(qn('w:rPr')).find(qn('w:b')) is not None \
            and re.match(r'^[A-G]\. ', txt):
        add_par('H2_Cont', text=txt); continue
    add_par('PARA', runs_from=el)

# AI-system reference and cross-reference numbers in the acknowledgment
ai_n = ref_n + 1
add_par('References', text=f'[{ai_n}] Anthropic, "Claude," AI system, 2026. [Online]. Available: https://www.anthropic.com/claude')
for p in doc.paragraphs:
    if '{ATLAS}' in p.text:
        for r in p.runs:
            r.text = r.text.replace('{ATLAS}', f'[{atlas_ref}]').replace('{AI}', f'[{ai_n}]')

# ---- 4. author biography ------------------------------------------------------------
p = add_par('AU_Bios')
r = p.add_run('MURAT ARSLAN '); r.bold = True
p.add_run('is an Assistant Professor with the Department of Software Engineering, Altınbaş Cyprus University, Cyprus. ')
r = p.add_run('[Degrees, previous positions, research interests and author photograph to be completed by the author.]')
r.font.highlight_color = 7

# schema order for pPr / tblPr / tcPr children
PPR = ['pStyle','keepNext','keepLines','pageBreakBefore','framePr','widowControl','numPr','suppressLineNumbers','pBdr','shd','tabs',
       'suppressAutoHyphens','kinsoku','wordWrap','overflowPunct','topLinePunct','autoSpaceDE','autoSpaceDN','bidi','adjustRightInd',
       'snapToGrid','spacing','ind','contextualSpacing','mirrorIndents','suppressOverlap','jc','textDirection','textAlignment',
       'textboxTightWrap','outlineLvl','divId','cnfStyle','rPr','sectPr','pPrChange']
TBL = ['tblStyle','tblpPr','tblOverlap','bidiVisual','tblStyleRowBandSize','tblStyleColBandSize','tblW','jc','tblCellSpacing','tblInd',
       'tblBorders','shd','tblLayout','tblCellMar','tblLook','tblCaption','tblDescription']
TC = ['cnfStyle','tcW','gridSpan','hMerge','vMerge','tcBorders','shd','noWrap','tcMar','textDirection','tcFitText','vAlign','hideMark']
def reorder(el, order):
    kids = list(el)
    for k in kids: el.remove(k)
    kids.sort(key=lambda k: order.index(k.tag.split('}')[1]) if k.tag.split('}')[1] in order else 99)
    for k in kids: el.append(k)
for pp in body.iter(qn('w:pPr')): reorder(pp, PPR)
for tp in body.iter(qn('w:tblPr')):
    for dup in tp.findall(qn('w:tblW'))[:-1]: tp.remove(dup)
    reorder(tp, TBL)
for tc in body.iter(qn('w:tcPr')): reorder(tc, TC)
doc.save(OUT)
print('saved', OUT, 'refs', ref_n, 'atlas_ref', atlas_ref, 'ai_ref', ai_n)
