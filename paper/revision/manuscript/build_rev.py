import re, subprocess, sys
from docx import Document
from docx.shared import Pt, RGBColor, Inches, Twips
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

B = {
 'porcheron': 'M. Porcheron, J. E. Fischer, S. Reeves, and S. Sharples, "Voice interfaces in everyday life," in *Proc. ACM CHI Conf. Human Factors Comput. Syst.*, 2018, pp. 1–12.',
 'clark': 'L. Clark et al., "The state of speech in HCI: Trends, themes and challenges," *Interact. Comput.*, vol. 31, no. 4, pp. 349–371, 2019.',
 'cassell': 'J. Cassell, J. Sullivan, S. Prevost, and E. Churchill, *Embodied Conversational Agents*. Cambridge, MA, USA: MIT Press, 2000.',
 'breazeal': 'C. Breazeal, "Toward sociable robots," *Robot. Auton. Syst.*, vol. 42, no. 3–4, pp. 167–175, 2003.',
 'li': 'J. Li, "The benefit of being physically present: A survey of experimental works comparing copresent robots, telepresent robots and virtual agents," *Int. J. Hum.-Comput. Stud.*, vol. 77, pp. 23–37, 2015.',
 'skantze': 'G. Skantze, "Turn-taking in conversational systems and human-robot interaction: A review," *Comput. Speech Lang.*, vol. 67, p. 101178, 2021.',
 'stivers': 'T. Stivers et al., "Universals and cultural variation in turn-taking in conversation," *Proc. Natl. Acad. Sci. USA*, vol. 106, no. 26, pp. 10587–10592, 2009.',
 'gpt4o': 'OpenAI, "GPT-4o system card," OpenAI, Tech. Rep., 2024. [Online]. Available: https://openai.com/index/gpt-4o-system-card/',
 'gemini': 'Google DeepMind, "Gemini 2.5: Pushing the frontier with advanced reasoning, multimodality, long context, and next generation agentic capabilities," Google, Tech. Rep., 2025. [Online]. Available: https://deepmind.google/technologies/gemini/',
 'jali': 'P. Edwards, C. Landreth, E. Fiume, and K. Singh, "JALI: An animator-centric viseme model for expressive lip synchronization," *ACM Trans. Graph.*, vol. 35, no. 4, pp. 1–11, 2016.',
 'mfa': 'M. McAuliffe, M. Socolof, S. Mihuc, M. Wagner, and M. Sonderegger, "Montreal Forced Aligner: Trainable text-speech alignment using Kaldi," in *Proc. Interspeech*, 2017, pp. 498–502.',
 'taylor': 'S. Taylor et al., "A deep learning approach for generalized speech animation," *ACM Trans. Graph.*, vol. 36, no. 4, pp. 1–11, 2017.',
 'karras': 'T. Karras, T. Aila, S. Laine, A. Herva, and J. Lehtinen, "Audio-driven facial animation by joint end-to-end learning of pose and emotion," *ACM Trans. Graph.*, vol. 36, no. 4, pp. 1–12, 2017.',
 'faceformer': 'Y. Fan, Z. Lin, J. Saito, W. Wang, and T. Komura, "FaceFormer: Speech-driven 3D facial animation with transformers," in *Proc. IEEE/CVF Conf. Comput. Vis. Pattern Recognit. (CVPR)*, 2022, pp. 18770–18780.',
 'a2f3d': 'NVIDIA, C. Chung et al., "Audio2Face-3D: Audio-driven realistic facial animation for digital avatars," *arXiv:2508.16401*, 2025.',
 'livatar': 'H. Liu et al., "Livatar-1: Real-time talking heads generation with tailored flow matching," *arXiv:2507.18649*, 2025.',
 'talkingmachines': 'C. Low and W. Wang, "TalkingMachines: Real-time audio-driven FaceTime-style video via autoregressive diffusion models," *arXiv:2506.03099*, 2025.',
 'ovrlipsync': 'Meta Platforms, "Oculus Lipsync SDK documentation," 2023. [Online]. Available: https://developers.meta.com/horizon/documentation/',
 'ulipsync': 'hecomi, "uLipSync: MFCC-based LipSync plug-in for Unity," GitHub repository, 2021. [Online]. Available: https://github.com/hecomi/uLipSync',
 'rhubarb': 'D. S. Wolf, "Rhubarb Lip Sync," GitHub repository, 2016. [Online]. Available: https://github.com/DanielSWolf/rhubarb-lip-sync',
 'agora': 'Agora Inc., "Build real-time AI avatars with lip sync using Agora ConvoAI," Technical Blog, 2025. [Online]. Available: https://www.agora.io/en/blog/build-real-time-ai-avatars-with-lip-sync-using-agora-convoai-rpm/',
 'mascot': 'Mascot Bot, "Gemini Live API avatar integration — interactive AI avatar with lip sync," SDK Documentation, 2025. [Online]. Available: https://docs.mascot.bot/libraries/gemini-live-api-avatar',
 'kiseleva': 'J. Kiseleva et al., "Understanding user satisfaction with intelligent assistants," in *Proc. ACM CHIIR*, 2016, pp. 121–130.',
 'lee': 'K. M. Lee, Y. Jung, J. Kim, and S. R. Kim, "Are physically embodied social agents better than disembodied social agents? The effects of physical embodiment, tactile interaction, and people\'s loneliness in human–robot interaction," *Int. J. Hum.-Comput. Stud.*, vol. 64, no. 10, pp. 962–973, 2006.',
 'vaswani': 'A. Vaswani et al., "Attention is all you need," in *Adv. Neural Inf. Process. Syst.*, vol. 30, 2017.',
 'brown': 'T. B. Brown et al., "Language models are few-shot learners," in *Adv. Neural Inf. Process. Syst.*, vol. 33, 2020, pp. 1877–1901.',
 'wei': 'J. Wei et al., "Chain-of-thought prompting elicits reasoning in large language models," in *Adv. Neural Inf. Process. Syst.*, vol. 35, 2022.',
 'toolformer': 'T. Schick et al., "Toolformer: Language models can teach themselves to use tools," in *Adv. Neural Inf. Process. Syst.*, vol. 36, 2023.',
 'park': 'J. S. Park, J. C. O\'Brien, C. J. Cai, M. R. Morris, P. Liang, and M. S. Bernstein, "Generative agents: Interactive simulacra of human behavior," in *Proc. ACM UIST*, 2023, pp. 1–22.',
 'fisher': 'C. G. Fisher, "Confusions among visually perceived consonants," *J. Speech Hear. Res.*, vol. 11, no. 4, pp. 796–804, 1968.',
 'cohen': 'M. M. Cohen and D. W. Massaro, "Modeling coarticulation in synthetic visual speech," in *Models and Techniques in Computer Animation*, N. M. Thalmann and D. Thalmann, Eds. Tokyo, Japan: Springer, 1993, pp. 139–156.',
 'whisper': 'A. Radford et al., "Robust speech recognition via large-scale weak supervision," in *Proc. ICML*, 2023, pp. 28492–28518.',
 'whisperx': 'M. Bain, J. Huh, T. Han, and A. Zisserman, "WhisperX: Time-accurate speech transcription of long-form audio," in *Proc. Interspeech*, 2023, pp. 4489–4493.',
 'ctcseg': 'L. Kürzinger, D. Winkelbauer, L. Li, T. Watzel, and G. Rigoll, "CTC-segmentation of large corpora for German end-to-end speech recognition," in *Proc. SPECOM*, 2020, pp. 267–278.',
 'voca': 'D. Cudeiro, T. Bolkart, C. Laidlaw, A. Ranjan, and M. J. Black, "Capture, learning, and synthesis of 3D speaking styles," in *Proc. IEEE/CVF CVPR*, 2019, pp. 10101–10111.',
 'meshtalk': 'A. Richard, M. Zollhöfer, Y. Wen, F. de la Torre, and Y. Sheikh, "MeshTalk: 3D face animation from speech using cross-modality disentanglement," in *Proc. IEEE/CVF ICCV*, 2021, pp. 1173–1182.',
 'codetalker': 'J. Xing, M. Xia, Y. Zhang, X. Cun, J. Wang, and T.-T. Wong, "CodeTalker: Speech-driven 3D facial animation with discrete motion prior," in *Proc. IEEE/CVF CVPR*, 2023, pp. 12780–12790.',
 'selftalk': 'Z. Peng et al., "SelfTalk: A self-supervised commutative training diagram to comprehend 3D talking faces," in *Proc. ACM Multimedia*, 2023, pp. 5292–5301.',
 'unitalker': 'X. Fan et al., "UniTalker: Scaling up audio-driven 3D facial animation through a unified model," in *Proc. ECCV*, 2024.',
 'wav2lip': 'K. R. Prajwal, R. Mukhopadhyay, V. P. Namboodiri, and C. V. Jawahar, "A lip sync expert is all you need for speech to lip generation in the wild," in *Proc. ACM Multimedia*, 2020, pp. 484–492.',
 'vrm': 'VRM Consortium, "VRM: A file format for 3D avatars," 2024. [Online]. Available: https://vrm.dev/en/',
 'threevrm': 'Pixiv Inc., "three-vrm: VRM utilities for three.js," 2024. [Online]. Available: https://github.com/pixiv/three-vrm',
 'arkit': 'Apple Inc., "ARFaceAnchor.BlendShapeLocation — ARKit developer documentation," 2023. [Online]. Available: https://developer.apple.com/documentation/arkit/',
 'whisperstreaming': 'D. Macháček, R. Dabre, and O. Bojar, "Turning Whisper into real-time transcription system," in *Proc. IJCNLP-AACL System Demonstrations*, 2023, pp. 17–24.',
 'syncnet': 'J. S. Chung and A. Zisserman, "Out of time: Automated lip sync in the wild," in *Proc. ACCV Workshops*, 2016, pp. 251–263.',
 'bt1359': 'International Telecommunication Union, "Relative timing of sound and vision for broadcasting," ITU-R Recommendation BT.1359-1, 1998.',
 'mediapipe': 'C. Lugaresi et al., "MediaPipe: A framework for building perception pipelines," *arXiv:1906.08172*, 2019.',
 'grid': 'M. Cooke, J. Barker, S. Cunningham, and X. Shao, "An audio-visual corpus for speech perception and automatic speech recognition," *J. Acoust. Soc. Am.*, vol. 120, no. 5, pp. 2421–2424, 2006.',
 'streamingtalker': 'Y. Yang et al., "StreamingTalker: Audio-driven 3D facial animation with autoregressive diffusion model," *arXiv:2511.14223*, 2025.',
 'telepresence': 'J. Lee et al., "Audio driven real-time facial animation for social telepresence," *arXiv:2510.01176*, 2025.',
 'echoavatar': 'B. Chen, Y. Li, Y. Xu, Y. Zheng, Y. Weng, and K. Zhou, "EchoAvatar: Real-time generative avatar animation from audio streams," in *Proc. ACM SIGGRAPH Conf. Papers*, 2026, doi: 10.1145/3799902.3811066.',
 'tinyv2f': 'Z. Han, M. Teye, D. Yadgaroff, and J. Bütepage, "Tiny is not small enough: High-quality, low-resource facial animation models through hybrid knowledge distillation," *ACM Trans. Graph.*, vol. 44, no. 4, 2025, doi: 10.1145/3730929.',
 'w2v2': 'A. Baevski, Y. Zhou, A. Mohamed, and M. Auli, "wav2vec 2.0: A framework for self-supervised learning of speech representations," in *Adv. Neural Inf. Process. Syst.*, vol. 33, 2020, pp. 12449–12460.',
 'mms': 'V. Pratap et al., "Scaling speech technology to 1,000+ languages," *J. Mach. Learn. Res.*, vol. 25, no. 97, pp. 1–52, 2024.',
 'atlasmk37': 'FatihMakes, "Atlas-MK37 (MARK XXXVII): Cross-platform personal AI assistant," GitHub repository, 2026. [Online]. Available: https://github.com/FatihMakes/Atlas-MK37',
 'fleurs': 'A. Conneau et al., "FLEURS: Few-shot learning evaluation of universal representations of speech," in *Proc. IEEE Spoken Language Technology Workshop (SLT)*, 2022, pp. 798–805.',
 'lrs3': 'T. Afouras, J. S. Chung, and A. Zisserman, "LRS3-TED: A large-scale dataset for visual speech recognition," *arXiv:1809.00496*, 2018.',
 'avhubert': 'B. Shi, W.-N. Hsu, K. Lakhotia, and A. Mohamed, "Learning audio-visual speech representation by masked multimodal cluster prediction," in *Proc. ICLR*, 2022.',
 'sus': 'J. Brooke, "SUS: A \'quick and dirty\' usability scale," in *Usability Evaluation in Industry*, P. W. Jordan et al., Eds. London, U.K.: Taylor & Francis, 1996, pp. 189–194.',
 'bangor': 'A. Bangor, P. T. Kortum, and J. T. Miller, "An empirical evaluation of the System Usability Scale," *Int. J. Hum.–Comput. Interact.*, vol. 24, no. 6, pp. 574–594, 2008.',
 'ues': 'H. L. O\'Brien, P. Cairns, and M. Hall, "A practical approach to measuring user engagement with the refined user engagement scale (UES) and new UES short form," *Int. J. Hum.-Comput. Stud.*, vol. 112, pp. 28–39, 2018.',
 'jian': 'J.-Y. Jian, A. M. Bisantz, and C. G. Drury, "Foundations for an empirically determined scale of trust in automated systems," *Int. J. Cogn. Ergon.*, vol. 4, no. 1, pp. 53–71, 2000.',
}

src = open('paper_rev.md').read()
order = []
def num(m):
    keys = [k.strip() for k in m.group(1).split(',')]
    ns = []
    for k in keys:
        if k not in B: sys.exit('missing ref ' + k)
        if k not in order: order.append(k)
        ns.append(order.index(k) + 1)
    ns = sorted(set(ns))
    # compress ranges: [1]–[3]
    out, i = [], 0
    while i < len(ns):
        j = i
        while j + 1 < len(ns) and ns[j + 1] == ns[j] + 1: j += 1
        out.append(f'\\[{ns[i]}\\]' if j == i else (f'\\[{ns[i]}\\], \\[{ns[j]}\\]' if j == i + 1 else f'\\[{ns[i]}\\]–\\[{ns[j]}\\]'))
        i = j + 1
    return ', '.join(out)
body, refs_marker = src.split('{{REFERENCES}}')
body = re.sub(r'\{\{([a-z0-9_, ]+)\}\}', num, body)
body = re.sub(r'(?<!\\)_', r'\\_', body)
unused = [k for k in B if k not in order]
if unused: print('unused refs:', unused)
refs = '\n\n'.join(f'::: {{custom-style="RefText"}}\n\\[{i+1}\\] {B[k]}\n:::' for i, k in enumerate(order))
open('paper_rev.num.md', 'w').write(body + refs + refs_marker)
print('refs:', len(order))

out = 'TranscriptionSync_revised.docx'
subprocess.run(['pandoc', 'paper_rev.num.md', '-f', 'markdown-smart-superscript-subscript', '-o', out], check=True)

d = Document(out)
sec = d.sections[0]
sec.page_width, sec.page_height = Twips(12240), Twips(15840)
for side in ('top_margin', 'bottom_margin', 'left_margin', 'right_margin'):
    setattr(sec, side, Twips(1247))
NAVY = RGBColor(0x1F, 0x35, 0x64)
st = d.styles
def font(s, size=None, bold=None, color=None, italic=None):
    f = s.font; f.name = 'Calibri'
    rpr = s.element.get_or_add_rPr(); rf = rpr.find(qn('w:rFonts'))
    if rf is None: rf = OxmlElement('w:rFonts'); rpr.append(rf)
    for a in ('w:ascii', 'w:hAnsi', 'w:cs', 'w:eastAsia'): rf.set(qn(a), 'Calibri')
    for a in ('w:asciiTheme', 'w:hAnsiTheme', 'w:cstheme', 'w:eastAsiaTheme'):
        if rf.get(qn(a)) is not None: del rf.attrib[qn(a)]
    if size: f.size = Pt(size)
    if bold is not None: f.bold = bold
    if italic is not None: f.italic = italic
    if color is not None: f.color.rgb = color
for name in ('Normal', 'Body Text', 'First Paragraph', 'Compact'):
    try:
        s = st[name]; font(s, 10.5, color=RGBColor(0, 0, 0))
        s.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        s.paragraph_format.space_after = Pt(6); s.paragraph_format.space_before = Pt(0)
        s.paragraph_format.line_spacing = 1.1
    except KeyError: pass
s = st['PaperTitle']; font(s, 15, True, NAVY); s.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER; s.paragraph_format.space_after = Pt(12)
s = st['SecHead']; font(s, 12.5, True, NAVY); s.paragraph_format.space_before = Pt(14); s.paragraph_format.space_after = Pt(6); s.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT; s.paragraph_format.keep_with_next = True
s = st['CaptionText']; font(s, 9, color=RGBColor(0, 0, 0)); s.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY; s.paragraph_format.space_after = Pt(8); s.paragraph_format.keep_with_next = False
s = st['Equation']; font(s, 10.5); s.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
s = st['RefText']; font(s, 9); s.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT; s.paragraph_format.space_after = Pt(3)
# figures centered
for p in d.paragraphs:
    if p._p.xpath('.//w:drawing'):
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.keep_with_next = True
# table captions keep with next (table)
body_el = d.element.body
for p in d.paragraphs:
    if p.style.name == 'CaptionText' and p.text.startswith('TABLE'):
        p.paragraph_format.keep_with_next = True
# tables: grid, header shading, widths
avail = 12240 - 2 * 1247
for t in d.tables:
    tbl = t._tbl; pr = tbl.tblPr
    for old in pr.findall(qn('w:tblBorders')): pr.remove(old)
    b = OxmlElement('w:tblBorders')
    for e in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        x = OxmlElement(f'w:{e}'); x.set(qn('w:val'), 'single'); x.set(qn('w:sz'), '4'); x.set(qn('w:color'), '808080'); b.append(x)
    pr.append(b)
    jc = pr.find(qn('w:jc'))
    if jc is None: jc = OxmlElement('w:jc'); pr.append(jc)
    jc.set(qn('w:val'), 'center')
    ncol = len(t.columns)
    # weight columns by max text length
    lens = [max(8, max(min(len(r.cells[i].text), 40) for r in t.rows)) for i in range(ncol)]
    tot = sum(lens); ws = [int(avail * l / tot) for l in lens]; ws[-1] = avail - sum(ws[:-1])
    tw = pr.find(qn('w:tblW'))
    if tw is None: tw = OxmlElement('w:tblW'); pr.append(tw)
    tw.set(qn('w:w'), str(avail)); tw.set(qn('w:type'), 'dxa')
    [pr.remove(x) for x in pr.findall(qn('w:tblLayout'))]
    lay = OxmlElement('w:tblLayout'); lay.set(qn('w:type'), 'fixed'); pr.append(lay)
    grid = tbl.tblGrid
    for i, gc in enumerate(grid.findall(qn('w:gridCol'))): gc.set(qn('w:w'), str(ws[i]))
    for ri, row in enumerate(t.rows):
        for ci, cell in enumerate(row.cells):
            tcpr = cell._tc.get_or_add_tcPr()
            w = tcpr.find(qn('w:tcW'))
            if w is None: w = OxmlElement('w:tcW'); tcpr.append(w)
            w.set(qn('w:w'), str(ws[ci])); w.set(qn('w:type'), 'dxa')
            if ri == 0:
                sh = OxmlElement('w:shd'); sh.set(qn('w:val'), 'clear'); sh.set(qn('w:color'), 'auto'); sh.set(qn('w:fill'), 'BDD7EE'); tcpr.append(sh)
            for p in cell.paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                p.paragraph_format.space_after = Pt(1); p.paragraph_format.space_before = Pt(1)
                for r in p.runs: r.font.size = Pt(8 if ncol >= 7 else 9)
TBL=['tblStyle','tblpPr','tblOverlap','bidiVisual','tblStyleRowBandSize','tblStyleColBandSize','tblW','jc','tblCellSpacing','tblInd','tblBorders','shd','tblLayout','tblCellMar','tblLook','tblCaption','tblDescription']
TC=['cnfStyle','tcW','gridSpan','hMerge','vMerge','tcBorders','shd','noWrap','tcMar','textDirection','tcFitText','vAlign','hideMark']
def reorder(el, order):
    kids=list(el); 
    for k in kids: el.remove(k)
    kids.sort(key=lambda k: order.index(k.tag.split('}')[1]) if k.tag.split('}')[1] in order else 99)
    for k in kids: el.append(k)
for t in d.tables:
    reorder(t._tbl.tblPr, TBL)
    for tc in t._tbl.iter(qn('w:tc')): reorder(tc.get_or_add_tcPr(), TC)
pm=d.element.body.find(qn('w:sectPr')).find(qn('w:pgMar'))
for a,v in (('w:header','720'),('w:footer','720'),('w:gutter','0')): pm.set(qn(a),v)
d.save(out)
print('saved', out)
